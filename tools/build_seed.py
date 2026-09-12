#!/usr/bin/env python3
"""Excel「見積実行予算作成ツール」から db/seed_real.sql を生成する。

入力: tools/source/見積実行予算作成ツール*.xlsx
出力: db/seed_real.sql（カテゴリ・単価マスタ・選択肢・特記事項ひな形・予算項目・管理者・サンプル案件）
"""
import glob
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sorted(glob.glob(os.path.join(ROOT, "tools/source/見積実行予算作成ツール*.xlsx")))
OUT = os.path.join(ROOT, "db/seed_real.sql")
ADMIN_PW = os.environ.get("KANAME_ADMIN_PW", "kaname2026")


def q(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:
            return "NULL"
        return repr(v) if isinstance(v, float) else str(v)
    s = str(v).replace("\\", "\\\\").replace("'", "''")
    return "'" + s + "'"


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def bcrypt(pw):
    import subprocess

    return subprocess.check_output(
        ["php", "-r", "echo password_hash($argv[1], PASSWORD_DEFAULT);", pw], text=True
    ).strip()


def main():
    if not SRC:
        print("tools/source にExcelがありません", file=sys.stderr)
        sys.exit(1)
    path = SRC[-1]
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    w = out.append
    w("-- 自動生成: tools/build_seed.py  元ファイル: " + os.path.basename(path))
    w("SET NAMES utf8mb4;")
    w("SET @now = NOW();")

    # ---------- 管理者 ----------
    w("INSERT INTO users (login_id, name, password_hash, role, is_active, must_change_pw, created_at, updated_at)")
    w(f"SELECT 'admin', '管理者', {q(bcrypt(ADMIN_PW))}, 'admin', 1, 0, @now, @now FROM DUAL WHERE NOT EXISTS (SELECT 1 FROM users WHERE login_id='admin');")

    # ---------- カテゴリ・品目（データ） ----------
    ws = wb["データ"]
    cats = []
    items = []
    cat_name = None
    seen = {}
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if isinstance(b, str) and b.strip().startswith("『"):
            cat_name = b.strip().strip("『』").replace("\u3000", "").strip()
            cats.append(cat_name)
            continue
        code = txt(a)
        if code is None or code == "リストへ":
            continue
        name = txt(b) or ""
        if code in seen:
            seen[code] += 1
            code = f"{code}-{seen[code]}"
        else:
            seen[code] = 1
        note = txt(ws.cell(r, 17).value)
        items.append(
            dict(
                code=code,
                cat=cat_name,
                sort=r,
                name=name,
                material=txt(ws.cell(r, 3).value),
                length=num(ws.cell(r, 4).value),
                width=num(ws.cell(r, 5).value),
                thickness=num(ws.cell(r, 6).value),
                area_unit=txt(ws.cell(r, 7).value),
                use_unit=txt(ws.cell(r, 8).value),
                material_price=num(ws.cell(r, 9).value),
                labor_rate=num(ws.cell(r, 10).value),
                labor_unit_price=num(ws.cell(r, 11).value),
                note=note,
            )
        )
    for i, c in enumerate(cats, 1):
        w(f"INSERT INTO categories (sort_no, name) VALUES ({i}, {q(c)}) ON DUPLICATE KEY UPDATE sort_no=VALUES(sort_no);")
    w("INSERT INTO items (code, category_id, sort_no, name, material, length, width, thickness, area_unit, use_unit, material_price, labor_rate, labor_unit_price, note, is_active, created_at, updated_at) VALUES")
    rows = []
    for it in items:
        cat_sql = f"(SELECT id FROM categories WHERE name={q(it['cat'])})" if it["cat"] else "NULL"
        rows.append(
            "(" + ", ".join(
                [
                    q(it["code"]), cat_sql, str(it["sort"]), q(it["name"]), q(it["material"]),
                    q(it["length"]), q(it["width"]), q(it["thickness"]), q(it["area_unit"]), q(it["use_unit"]),
                    q(it["material_price"]), q(it["labor_rate"]), q(it["labor_unit_price"]), q(it["note"]), "1", "@now", "@now",
                ]
            ) + ")"
        )
    w(",\n".join(rows))
    w("ON DUPLICATE KEY UPDATE category_id=VALUES(category_id), sort_no=VALUES(sort_no), name=VALUES(name), material=VALUES(material), length=VALUES(length), width=VALUES(width), thickness=VALUES(thickness), area_unit=VALUES(area_unit), use_unit=VALUES(use_unit), material_price=VALUES(material_price), labor_rate=VALUES(labor_rate), labor_unit_price=VALUES(labor_unit_price), note=VALUES(note), updated_at=@now;")

    # ---------- 選択肢（経理実行予算 59行〜） ----------
    ws = wb["経理実行予算"]

    def col_values(col, r0=61, r1=96):
        vals = []
        for r in range(r0, r1 + 1):
            v = ws.cell(r, col).value
            if v in (None, ""):
                continue
            if isinstance(v, str) and v.startswith("【"):
                break
            vals.append(v)
        return vals

    def colnum(letter):
        n = 0
        for ch in letter:
            n = n * 26 + (ord(ch) - 64)
        return n

    lists = {
        "pay_close": col_values(colnum("D"), 61, 66),
        "pay_day": col_values(colnum("E"), 61, 66),
        "pay_cash_pct": col_values(colnum("F"), 61, 66),
        "pay_bill_pct": col_values(colnum("G"), 61, 66),
        "pay_site_days": col_values(colnum("H"), 61, 66),
        "pay_terms": col_values(colnum("D"), 69, 75),
        "roof_spec": col_values(colnum("J"), 61, 89),
        "roof_material": col_values(colnum("L"), 61, 76),
        "roof_thickness": col_values(colnum("N"), 61, 73),
        "roof_product": col_values(colnum("P"), 61, 74),
        "info_source": col_values(colnum("S"), 61, 80),
        "repeat_kind": col_values(colnum("V"), 61, 63),
        "work_type": col_values(colnum("Y"), 61, 66),
        "receipt": col_values(colnum("AE"), 61, 64),
        "billing": col_values(colnum("AG"), 61, 64),
        "accounting": col_values(colnum("AE"), 70, 75),
    }
    w("DELETE FROM option_lists WHERE list_key IN (" + ",".join(q(k) for k in list(lists) + ["bank"]) + ");")
    for key, vals in lists.items():
        for i, v in enumerate(vals, 1):
            if isinstance(v, float) and v == int(v):
                v = int(v)
            w(f"INSERT INTO option_lists (list_key, sort_no, value) VALUES ({q(key)}, {i}, {q(str(v))});")
    # 銀行手数料率（営業実行予算 J3:L7）
    wsb = wb["営業実行予算"]
    i = 0
    for r in range(3, 8):
        name = txt(wsb.cell(r, 10).value)
        if not name:
            continue
        i += 1
        extra = f"{wsb.cell(r, 11).value},{wsb.cell(r, 12).value}"
        w(f"INSERT INTO option_lists (list_key, sort_no, value, extra) VALUES ('bank', {i}, {q(name)}, {q(extra)});")

    # ---------- 特記事項ひな形 ----------
    wsn = wb["工事全般特記事項"]
    notes = []
    for r in range(1, wsn.max_row + 1):
        a = wsn.cell(r, 1).value
        b = wsn.cell(r, 2).value
        if isinstance(a, (int, float)) and txt(b):
            notes.append((int(a), str(b).strip()))
    w("DELETE FROM note_templates;")
    for n, body in notes:
        w(f"INSERT INTO note_templates (sort_no, body) VALUES ({n}, {q(body)});")

    # ---------- 予算項目 ----------
    budget = []
    for r in range(6, 25):
        v = txt(ws.cell(r, 4).value)
        if v:
            budget.append(("material", r, v))
    for r in range(26, 55):
        v = txt(ws.cell(r, 4).value)
        if v:
            budget.append(("subcontract", r, v))
    for r in range(6, 14):
        v = txt(ws.cell(r, 13).value)
        if v:
            budget.append(("expense", r, v))
    w("INSERT INTO budget_items (group_key, sort_no, name) SELECT g, s, n FROM (SELECT NULL g, NULL s, NULL n FROM DUAL WHERE 0")
    for g, s, n in budget:
        w(f"  UNION ALL SELECT {q(g)}, {s}, {q(n)}")
    w(") t WHERE NOT EXISTS (SELECT 1 FROM budget_items b WHERE b.group_key=t.g AND b.name=t.n AND b.sort_no=t.s);")

    # ---------- サンプル案件（明細書・経理実行予算の現在値） ----------
    wsm = wb["明細書"]
    pname = txt(wsm.cell(3, 3).value) or "サンプル案件"
    w(f"SET @exists = (SELECT COUNT(*) FROM projects WHERE name={q(pname)});")
    w("SET @skip = @exists > 0;")

    def cell(sheet, ref):
        return sheet[ref].value

    def dt(v):
        if v is None or isinstance(v, str):
            return None
        return v.strftime("%Y-%m-%d")

    pcols = {
        "code": q(txt(cell(ws, "B1"))),
        "name": q(pname),
        "staff_name": q(txt(cell(ws, "W22"))),
        "site_name": q(txt(cell(ws, "O22"))),
        "site_zip": q(txt(cell(ws, "P23"))),
        "site_address": q(txt(cell(ws, "S23"))),
        "site_tel": q(txt(cell(ws, "Q24"))),
        "site_fax": q(txt(cell(ws, "U24"))),
        "period_from": q(dt(cell(ws, "O25"))),
        "period_to": q(dt(cell(ws, "S25"))),
        "start_date": q(dt(cell(ws, "O25"))),
        "pay_terms_text": q(txt(cell(wsb, "G22"))),
        "uniform_labor_price": q(num(cell(wsm, "P3"))),
        "uniform_rate": q(num(cell(wsm, "X3")) or 0.8),
        "general_admin_rate": q(num(cell(wsb, "G18")) or 0.08),
        "site_admin_rate": q(num(cell(wsb, "G19")) or 0.1),
        "fee_rate": q((num(cell(wsb, "C9")) or 3) / 100),
        "planned_contract_amount": q(num(cell(wsb, "B24"))),
    }
    w("INSERT INTO projects (" + ", ".join(pcols) + ", status, created_at, updated_at)")
    w("SELECT " + ", ".join(pcols.values()) + ", 'draft', @now, @now FROM DUAL WHERE NOT @skip;")
    w("SET @pid = IF(@skip, NULL, LAST_INSERT_ID());")

    line_key_re = re.compile(r"^(\d+)行目$")
    zen = str.maketrans("０１２３４５６７８９", "0123456789")
    block_starts = [5 + 31 * i for i in range(15)]
    for si, start in enumerate(block_starts, 1):
        sname = txt(wsm.cell(start, 3).value)
        if not sname:
            continue
        loss = num(wsm.cell(start, 8).value)
        loss = 0.07 if loss is None else loss
        w(f"INSERT INTO project_sections (project_id, sort_no, name, loss_rate) SELECT @pid, {si}, {q(sname)}, {loss} FROM DUAL WHERE NOT @skip;")
        w("SET @sid = IF(@skip, NULL, LAST_INSERT_ID());")
        lines = []
        for r in range(start + 2, start + 29):
            code = txt(wsm.cell(r, 2).value)
            qty = num(wsm.cell(r, 8).value)
            rkey = txt(wsm.cell(r, 18).value)
            tkey = txt(wsm.cell(r, 20).value)
            rate = num(wsm.cell(r, 24).value)
            pname_edit = txt(wsm.cell(r, 28).value)
            pmat_edit = txt(wsm.cell(r, 30).value)
            remarks = txt(wsm.cell(r, 38).value)
            if code is None and qty is None and not pname_edit:
                continue

            def keyno(k):
                if not k:
                    return None
                m = line_key_re.match(k.translate(zen))
                return int(m.group(1)) if m else None

            lines.append(dict(row=r, code=code, qty=qty, rkey=keyno(rkey), tkey=keyno(tkey), rate=rate,
                              pname=pname_edit, pmat=pmat_edit, remarks=remarks))
        # 行キー → 行インデックス
        keymap = {ln["rkey"]: idx for idx, ln in enumerate(lines) if ln["rkey"]}
        merges = []
        for idx, ln in enumerate(lines, 1):
            is_quote = 1 if ln["rkey"] else 0
            merge_idx = keymap.get(ln["tkey"]) if (not ln["rkey"] and ln["tkey"]) else None
            if merge_idx is not None:
                merges.append((idx, merge_idx + 1))
            code_sql = q(ln["code"])
            item_sel = f"(SELECT id FROM items WHERE code={code_sql})" if ln["code"] else "NULL"
            w("INSERT INTO project_lines (section_id, sort_no, item_id, item_code, name, material, length, width, thickness, area_unit, use_unit, material_price, labor_rate, labor_unit_price, quantity, is_quote, merge_into_line_id, rate_override, print_name, print_material, remarks)")
            w(f"SELECT @sid, {idx}, i.id, {code_sql}, i.name, i.material, i.length, i.width, i.thickness, i.area_unit, i.use_unit, i.material_price, i.labor_rate, i.labor_unit_price, {q(ln['qty'])}, {is_quote}, NULL, {q(ln['rate'])}, {q(ln['pname'])}, {q(ln['pmat'])}, {q(ln['remarks'])}")
            w(f"FROM (SELECT 1) d LEFT JOIN items i ON i.id = {item_sel} WHERE NOT @skip;")
        for idx, target_no in merges:
            w(f"UPDATE project_lines SET merge_into_line_id = (SELECT id FROM (SELECT id FROM project_lines WHERE section_id=@sid AND sort_no={target_no}) t) WHERE section_id=@sid AND sort_no={idx} AND NOT @skip;")

    # 特記事項（案件へコピー）
    for n, body in notes:
        w(f"INSERT INTO project_notes (project_id, sort_no, body) SELECT @pid, {n}, {q(body)} FROM DUAL WHERE NOT @skip;")
    # その他工事
    for no in range(1, 5):
        w(f"INSERT INTO extra_works (project_id, no, title) SELECT @pid, {no}, {q('その他工事' + str(no))} FROM DUAL WHERE NOT @skip;")

    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")
    print(f"categories={len(cats)} items={len(items)} notes={len(notes)} budget_items={len(budget)} -> {OUT}")


if __name__ == "__main__":
    main()
